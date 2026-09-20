import docx
from typing import List
from backend.app.loaders.base_loader import BaseDocumentLoader, DocumentContent
from backend.app.utils.text_cleaner import clean_extracted_text

class DOCXLoader(BaseDocumentLoader):
    def load(self, file_path: str, filename: str) -> List[DocumentContent]:
        contents: List[DocumentContent] = []
        try:
            doc = docx.Document(file_path)
            current_section = "Introduction"
            section_buffer = []

            # Process paragraphs and detect headings
            for p in doc.paragraphs:
                text = p.text.strip()
                if not text:
                    continue

                style_name = p.style.name.lower() if p.style and p.style.name else ""
                if "heading" in style_name or p.style.name.startswith("Heading"):
                    # Flush previous section buffer
                    if section_buffer:
                        full_section_text = clean_extracted_text("\n".join(section_buffer))
                        if full_section_text:
                            contents.append(
                                DocumentContent(
                                    text=full_section_text,
                                    metadata={
                                        "filename": filename,
                                        "source": filename,
                                        "section": current_section,
                                    }
                                )
                            )
                        section_buffer = []
                    current_section = text

                section_buffer.append(text)

            # Process tables in DOCX
            for table_idx, table in enumerate(doc.tables):
                table_lines = []
                headers = []
                for row_idx, row in enumerate(table.rows):
                    cells = [cell.text.strip().replace("\n", " ") for cell in row.cells]
                    # Avoid duplicated adjacent merged cells
                    cleaned_cells = []
                    for i, c in enumerate(cells):
                        if i == 0 or c != cells[i - 1]:
                            cleaned_cells.append(c)

                    if row_idx == 0:
                        headers = cleaned_cells
                        table_lines.append(" | ".join(headers))
                        table_lines.append("-" * max(20, len(" | ".join(headers))))
                    else:
                        table_lines.append(" | ".join(cleaned_cells))

                if table_lines:
                    table_text = clean_extracted_text("\n".join(table_lines))
                    if table_text:
                        contents.append(
                            DocumentContent(
                                text=f"[Table {table_idx + 1}]\n{table_text}",
                                metadata={
                                    "filename": filename,
                                    "source": filename,
                                    "section": f"Table {table_idx + 1}",
                                }
                            )
                        )

            # Flush remaining section buffer
            if section_buffer:
                full_section_text = clean_extracted_text("\n".join(section_buffer))
                if full_section_text:
                    contents.append(
                        DocumentContent(
                            text=full_section_text,
                            metadata={
                                "filename": filename,
                                "source": filename,
                                "section": current_section,
                            }
                        )
                    )

            if not contents:
                raise ValueError("DOCX file contains no extractable text or tables.")

            return contents

        except Exception as e:
            raise ValueError(f"Failed to process DOCX file '{filename}': {str(e)}")
